import re
import sys
from youtube_transcript_api import YouTubeTranscriptApi, TranscriptsDisabled, NoTranscriptFound
from docx import Document
from urllib.parse import urlparse, parse_qs

def extract_video_id(url):
    """Extracts the YouTube video ID from various URL formats."""
    # Examples:
    # https://www.youtube.com/watch?v=VIDEO_ID
    # https://youtu.be/VIDEO_ID
    # https://www.youtube.com/embed/VIDEO_ID
    # https://www.youtube.com/v/VIDEO_ID
    # https://m.youtube.com/watch?v=VIDEO_ID

    parsed_url = urlparse(url)

    # Standard youtube.com/watch?v=...
    if parsed_url.netloc in ('www.youtube.com', 'youtube.com', 'm.youtube.com'):
        if parsed_url.path == '/watch':
            query_params = parse_qs(parsed_url.query)
            if 'v' in query_params and len(query_params['v']) > 0:
                return query_params['v'][0]
        # Handle embed/v paths
        elif parsed_url.path.startswith(('/embed/', '/v/')):
             # Path is like /embed/VIDEO_ID or /v/VIDEO_ID
             path_parts = parsed_url.path.split('/')
             if len(path_parts) > 2 and path_parts[2]:
                 return path_parts[2]

    # Shortened youtu.be/VIDEO_ID
    elif parsed_url.netloc == 'youtu.be':
        # Path is like /VIDEO_ID
        path_parts = parsed_url.path.split('/')
        if len(path_parts) > 1 and path_parts[1]:
            return path_parts[1]

    # Fallback regex (less reliable but catches some edge cases)
    # This regex looks for standard 11-character YouTube IDs
    match = re.search(r'[a-zA-Z0-9_-]{11}', url)
    if match:
        # Basic check to avoid matching random strings in the URL
        potential_id = match.group(0)
        if len(potential_id) == 11:
             # Check if it's preceded by common delimiters
             preceded_by = url[match.start()-1:match.start()] if match.start() > 0 else ''
             if preceded_by in ['=', '/', '?', '&', ' ']:
                 return potential_id

    print(f"Error: Could not extract video ID from URL: {url}")
    return None

def get_transcript(video_id):
    """Fetches the transcript for a given video ID."""
    try:
        # Fetch available transcripts
        transcript_list = YouTubeTranscriptApi.list_transcripts(video_id)

        # Try to get manually created transcript first (often better quality)
        try:
            # Prioritize English, add more languages like ['en', 'es'] if needed
            transcript = transcript_list.find_manually_created_transcript(['en'])
            print("Found manually created transcript.")
        except NoTranscriptFound:
            # If no manual one, try finding a generated one
            try:
                # Prioritize English
                transcript = transcript_list.find_generated_transcript(['en'])
                print("Found auto-generated transcript.")
            except NoTranscriptFound:
                print(f"Error: No English transcript found for video ID {video_id}.")
                # Optionally list available languages
                available_langs = [t.language for t in transcript_list]
                if available_langs:
                     print(f"Available languages: {', '.join(available_langs)}")
                else:
                     print("No transcripts seem to be available in any language.")
                return None

        # Fetch the actual transcript data (list of dictionaries or objects)
        transcript_data = transcript.fetch()
        return transcript_data

    except TranscriptsDisabled:
        print(f"Error: Transcripts are disabled for video ID {video_id}.")
        return None
    except NoTranscriptFound:
         # This might catch cases where even listing fails, e.g., private video
         print(f"Error: No transcripts could be found at all for video ID {video_id}.")
         return None
    except Exception as e:
        print(f"An unexpected error occurred while fetching transcript: {e}")
        return None

def save_transcript_to_word(transcript_data, filename="transcript.docx"):
    """Saves the transcript data to a Word (.docx) file."""
    if not transcript_data:
        print("No transcript data to save.")
        return

    doc = Document()
    doc.add_heading('YouTube Video Transcript', 0) # Add a title

    full_text = ""
    for entry in transcript_data:
        # --- CORE FIX ---
        # Use attribute access (.text) as the error indicated 'entry' is an object
        try:
            text = entry.text
        except AttributeError:
            # Fallback in case it's unexpectedly a dict or malformed object
            print(f"Warning: Could not get 'text' attribute from entry: {entry}. Trying dictionary access.")
            try:
                 text = entry['text'] # Try dictionary access as a last resort
            except (KeyError, TypeError):
                 print(f"Error: Could not extract text from entry: {entry}. Skipping this entry.")
                 text = "" # Assign empty string to prevent errors below
        # --- END FIX ---

        full_text += text + " " # Add text segment followed by a space

    # Add the full transcript as one paragraph
    # Using strip() to remove any leading/trailing whitespace from the combined text
    paragraph = doc.add_paragraph(full_text.strip())

    # --- Optional: Add each segment as a new paragraph ---
    # for entry in transcript_data:
    #     try:
    #         text = entry.text
    #         doc.add_paragraph(text)
    #     except AttributeError:
    #          print(f"Warning: Could not get 'text' attribute from entry: {entry}. Skipping.")
    #     except Exception as e_para: # Catch other potential errors during paragraph add
    #         print(f"Error adding paragraph for entry {entry}: {e_para}")
    # --------------------------------------------------------

    try:
        doc.save(filename)
        print(f"Transcript successfully saved to: {filename}")
    except Exception as e:
        # Provide more context on save failure if possible
        print(f"Error saving Word document '{filename}': {e}")
        print("Please ensure you have write permissions in the directory and the file is not open elsewhere.")


def main():
    youtube_url = input("Enter the YouTube video URL: ")

    video_id = extract_video_id(youtube_url)

    if not video_id:
        sys.exit(1) # Exit if video ID extraction failed

    print(f"Extracted Video ID: {video_id}")
    print("Fetching transcript...")

    transcript = get_transcript(video_id)

    if transcript:
        # Create a filename suggestion based on video ID
        # Replace characters potentially invalid in filenames (optional but safer)
        safe_video_id = re.sub(r'[\\/*?:"<>|]',"_", video_id)
        output_filename = f"{safe_video_id}_transcript.docx"

        print(f"Saving transcript to {output_filename}...")
        save_transcript_to_word(transcript, output_filename)
    else:
        print("Could not retrieve or process transcript.")
        sys.exit(1)

if __name__ == "__main__":
    main()