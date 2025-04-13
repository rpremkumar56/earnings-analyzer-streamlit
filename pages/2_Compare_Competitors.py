# pages/2_Compare_Competitors.py

import streamlit as st
import os
import sys
from datetime import datetime, date
import pandas as pd
import traceback

# --- Ensure Modules are Found (if necessary) ---
# script_dir = os.path.dirname(os.path.abspath(__file__))
# project_root = os.path.abspath(os.path.join(script_dir, '..'))
# sys.path.append(project_root)

# --- Load Environment Variables & Configure API ---
from dotenv import load_dotenv
load_dotenv()

import google.generativeai as genai
try:
    api_key = os.getenv("GOOGLE_API_KEY")
    if not api_key:
        st.error("🔴 CRITICAL ERROR: GOOGLE_API_KEY not found.")
        st.stop()
    genai.configure(api_key=api_key)
    # print("Gemini API Key configured for Compare Competitors page.")
except Exception as e:
    st.error(f"🔴 CRITICAL ERROR: Failed to configure Google Gemini API: {e}")
    st.stop()

# --- Import Core Logic Modules ---
try:
    # Import necessary DB and Report functions
    from database_manager import get_db_connection, get_comparison_data
    # Import report generation helpers and Gemini caller from report_generator
    # (Alternatively, create a shared utils file)
    from report_generator import (
        add_heading as rg_add_heading,
        add_paragraph as rg_add_paragraph,
        set_col_widths as rg_set_col_widths,
        generate_gemini_content # Need this for qualitative comparison
    )
    from docx import Document # Need this to create the document object
    from docx.shared import Inches # For column widths
    from docx.enum.text import WD_PARAGRAPH_ALIGNMENT # For table formatting

except ImportError as e:
     st.error(f"🔴 Error importing backend modules: {e}. Check file paths.")
     st.stop()

# --- Constants ---
DEFAULT_OUTPUT_DIR = "output_reports"
DEFAULT_DB_PATH = None

# --- Page Specific Session State ---
if 'compare_results' not in st.session_state:
    st.session_state.compare_results = None
if 'compare_error' not in st.session_state:
    st.session_state.compare_error = None
if 'compare_report_path' not in st.session_state:
    st.session_state.compare_report_path = None
if 'compare_show_results' not in st.session_state:
    st.session_state.compare_show_results = False

# --- Comparison Logic Functions (Adapted from compare_competitors.py script) ---

def create_financial_comparison_table_data(comparison_data, primary_ticker, competitors):
    """Prepares DataFrame for the financial comparison table."""
    print("Preparing financial comparison table data...")
    metrics_to_compare = { # Define metrics and display names
        "Revenue ($M) QX (Table)": "Revenue ($M)",
        "Operating Margin (%) QX (Calc)": "Operating Margin (%)",
        "Net Margin (%) QX (Calc)": "Net Margin (%)",
        "Net Income ($M) QX (Table)": "Net Income ($M)",
        "Basic EPS ($) QX (Table)": "Basic EPS ($)",
        "TCV ($B) QX": "TCV ($B)", # Check TCV keys based on pdf_processor output
        "TCV ($M) QX": "TCV ($M)",
        "YoY Revenue Growth (%) QX": "YoY Revenue Growth (%)", # Add growth if available
        "CC Revenue Growth (%) QX": "CC Revenue Growth (%)",
    }
    tickers_in_order = [primary_ticker.upper()] + [c.upper() for c in competitors]
    table_data = []
    df_columns = ["Metric"] + tickers_in_order

    for metric_key_base, display_name in metrics_to_compare.items():
        row_data = {"Metric": display_name}
        found_metric_for_any = False
        for ticker in tickers_in_order:
            # Try variations of keys (Table, Calc, Headline, $M, $B)
            raw_value = "N/A"
            potential_keys = [
                metric_key_base,
                metric_key_base.replace("(Table)","(Calc)"),
                metric_key_base.replace("(Table)","(Headline)"),
                metric_key_base.replace("($M)","($B)"),
            ]
            for p_key in potential_keys:
                val = comparison_data.get(ticker, {}).get("financials", {}).get(p_key, {}).get('raw')
                if val is not None:
                    raw_value = val
                    found_metric_for_any = True # Mark if found for at least one company
                    break # Found value for this ticker
            row_data[ticker] = raw_value
        # Only add row if the metric was found for at least one company
        if found_metric_for_any:
            table_data.append(row_data)

    if not table_data:
        print("  No comparable financial data found for table.")
        return None

    df = pd.DataFrame(table_data, columns=df_columns)
    return df


def generate_qualitative_comparison_text(comparison_data, primary_ticker, competitors, theme):
    """Generates comparative text for a specific theme using Gemini."""
    print(f"Generating qualitative comparison for theme: {theme}...")
    tickers_in_order = [primary_ticker.upper()] + [c.upper() for c in competitors]
    theme_summaries = {}
    prompt_sections = [f"Compare and contrast the following summaries regarding '{theme}' for the specified companies for the analyzed period."]
    prompt_sections.append("Focus on similarities, differences in strategy, key initiatives, reported progress, and overall sentiment where available. Provide a concise comparative analysis.\n")
    found_data_count = 0

    for ticker in tickers_in_order:
        summary_data = comparison_data.get(ticker, {}).get("summaries", {}).get(theme)
        if summary_data and (summary_data.get('text') or summary_data.get('sentiment')):
            summary_text = f"Sentiment: {summary_data.get('sentiment', 'N/A')}\nSummary: {summary_data.get('text', '(No summary text)')}"
            theme_summaries[ticker] = summary_text
            prompt_sections.append(f"--- {ticker} ---")
            prompt_sections.append(summary_text)
            prompt_sections.append("-" * (len(ticker) + 8) + "\n")
            found_data_count += 1
        #else: No need to add placeholder text to the prompt

    if found_data_count < 2: # Need at least two companies to compare
        print(f"  Insufficient data (<2 companies) found for theme '{theme}'. Skipping comparison.")
        return f"Insufficient data found (less than 2 companies had summaries) to generate a comparison for '{theme}' for this period."

    prompt_sections.append("--- Comparison Analysis ---")
    prompt = "\n".join(prompt_sections)

    # Use the imported generate_gemini_content
    comparison_text = generate_gemini_content(prompt, context=f"Qualitative Comparison - {theme}")

    if comparison_text.startswith("Error:"):
        print(f"  Comparison generation failed for theme {theme}: {comparison_text}")
        return f"Comparison for '{theme}' could not be generated due to an API or processing error."
    else:
        print(f"  Comparison generated successfully for theme: {theme}")
        return comparison_text


def run_comparison_analysis(primary_ticker, competitors_list, year, quarter, db_path=None):
    """Fetches data and generates comparison results dict."""
    st.session_state.compare_results = None
    st.session_state.compare_error = None
    st.session_state.compare_report_path = None
    st.session_state.compare_show_results = False
    db_conn = None
    results = {"error": None, "financial_df": None, "qualitative": {}}

    try:
        st.info("Connecting to database...")
        db_conn = get_db_connection(db_path)
        all_tickers = [primary_ticker.upper()] + [c.upper() for c in competitors_list]

        st.info(f"Fetching data for {', '.join(all_tickers)} for FY{year} Q{quarter}...")
        comparison_data = get_comparison_data(db_conn, all_tickers, year, quarter)

        # Check if primary company data exists
        if not comparison_data.get(primary_ticker.upper(), {}).get("report_id"):
            raise ValueError(f"Data for primary ticker {primary_ticker.upper()} not found in database for FY{year} Q{quarter}. Please analyze it first using the 'Analyze Company' page.")

        st.info("Preparing financial comparison table...")
        financial_df = create_financial_comparison_table_data(comparison_data, primary_ticker, competitors_list)
        results["financial_df"] = financial_df # Store DataFrame or None

        st.info("Generating qualitative comparisons using Gemini...")
        # Define themes for qualitative comparison
        themes_to_compare = [
            "AI/GenAI", "Demand Environment/Pipeline", "Margins/Profitability/Costs",
            # Add more specific themes if needed and available in summaries
            # "Vertical Performance (e.g., BFSI, Retail)",
            # "Geographic Performance"
        ]
        for theme in themes_to_compare:
            with st.spinner(f"Generating comparison for {theme}..."):
                 qual_text = generate_qualitative_comparison_text(comparison_data, primary_ticker, competitors_list, theme)
                 results["qualitative"][theme] = qual_text

        st.session_state.compare_results = results
        st.session_state.compare_show_results = True
        st.success("Comparison analysis generated successfully!")

    except Exception as e:
        st.session_state.compare_error = f"Error during comparison: {e}\n{traceback.format_exc()}"
        st.error(st.session_state.compare_error)
    finally:
        if db_conn: db_conn.close(); print("Database connection closed.")

# --- Generate DOCX Report for Comparison ---
def generate_comparison_docx(primary_ticker, competitors, year, quarter, results, output_filename):
    """Generates the comparison DOCX report."""
    try:
        doc = Document()
        comp_string = ", ".join(competitors)
        rg_add_heading(doc, f"Competitor Analysis: FY{year} Q{quarter}", level=0)
        rg_add_paragraph(doc, f"Primary Company: {primary_ticker.upper()}")
        rg_add_paragraph(doc, f"Competitors Compared: {comp_string.upper()}")
        rg_add_paragraph(doc, f"Generated: {datetime.now().strftime('%Y-%m-%d %H:%M:%S')}")
        doc.add_paragraph()

        rg_add_heading(doc, "Financial Benchmark", level=1)
        df = results.get("financial_df")
        if df is not None and not df.empty:
            # Add table to document (Copied/adapted from compare_competitors.py)
            num_rows, num_cols = df.shape[0] + 1, df.shape[1]
            table = doc.add_table(rows=num_rows, cols=num_cols, style='Table Grid')
            table.autofit = False; table.allow_autofit = False
            metric_width = 2.5; competitor_width = max(1.2, (6.5 - metric_width) / len([primary_ticker] + competitors))
            col_widths = [metric_width] + [competitor_width] * (num_cols -1)
            rg_set_col_widths(table, col_widths)
            # Headers
            for j, col_name in enumerate(df.columns):
                 p = table.rows[0].cells[j].paragraphs[0]
                 run = p.add_run(col_name); run.font.bold = True
                 p.alignment = WD_PARAGRAPH_ALIGNMENT.CENTER if j > 0 else WD_PARAGRAPH_ALIGNMENT.LEFT
            # Data
            for i in range(df.shape[0]):
                for j, col_name in enumerate(df.columns):
                    value = str(df.iat[i, j])
                    p = table.rows[i + 1].cells[j].paragraphs[0]; p.add_run(value)
                    if j > 0: p.alignment = WD_PARAGRAPH_ALIGNMENT.RIGHT
        else:
            rg_add_paragraph(doc, "(No comparable financial data found)", italic=True)
        doc.add_paragraph()

        rg_add_heading(doc, "Qualitative Comparison", level=1)
        qualitative_results = results.get("qualitative", {})
        if qualitative_results:
             for theme, comparison_text in qualitative_results.items():
                 rg_add_heading(doc, theme, level=2)
                 rg_add_paragraph(doc, comparison_text)
                 doc.add_paragraph()
        else:
             rg_add_paragraph(doc, "(No qualitative comparisons generated)", italic=True)

        os.makedirs(os.path.dirname(output_filename), exist_ok=True)
        doc.save(output_filename)
        return True
    except Exception as e:
        st.session_state.compare_error = f"Error generating comparison DOCX: {e}\n{traceback.format_exc()}"
        st.error(st.session_state.compare_error)
        return False

# --- Streamlit UI ---
st.header("⚔️ Compare Competitors")
st.markdown("Select companies and a period to compare using data stored in the database.")

with st.form("compare_form"):
    st.subheader("Inputs:")
    col1, col2 = st.columns(2)
    with col1:
        primary_ticker = st.text_input("Primary Ticker Symbol", key="compare_primary")
        year = st.number_input("Fiscal Year", min_value=2000, max_value=date.today().year + 5, step=1, key="compare_year")
    with col2:
        # Allow comma-separated input for competitors
        competitors_str = st.text_input("Competitor Tickers (comma-separated)", key="compare_competitors")
        quarter = st.selectbox("Fiscal Quarter", options=[1, 2, 3, 4], key="compare_quarter")

    submitted = st.form_submit_button("📊 Generate Comparison")

    if submitted:
        st.session_state.compare_show_results = False # Reset display trigger
        st.session_state.compare_report_path = None # Reset report path
        if not primary_ticker or not competitors_str or not year or not quarter:
             st.warning("⚠️ Please provide Primary Ticker, Competitor Tickers, Year, and Quarter.")
        else:
             competitors_list = [c.strip() for c in competitors_str.split(',') if c.strip()]
             if not competitors_list:
                 st.warning("⚠️ Please provide at least one valid Competitor Ticker.")
             else:
                 with st.spinner("⏳ Fetching data and generating comparison... This may take time."):
                     run_comparison_analysis(primary_ticker, competitors_list, year, quarter, db_path=DEFAULT_DB_PATH)


# --- Display Results ---
if st.session_state.compare_show_results:
    if st.session_state.compare_results:
        st.divider()
        st.header("📈 Comparison Results")
        results = st.session_state.compare_results

        st.subheader("Financial Benchmark")
        if results.get("financial_df") is not None and not results["financial_df"].empty:
            st.dataframe(results["financial_df"].set_index("Metric"), use_container_width=True)
        else:
            st.info("No comparable financial data found for the selected period/companies.")

        st.subheader("Qualitative Comparison")
        qualitative_results = results.get("qualitative", {})
        if qualitative_results:
            for theme, comparison_text in qualitative_results.items():
                with st.expander(f"**{theme}** Comparison"):
                    st.markdown(comparison_text)
        else:
            st.info("No qualitative comparison summaries were generated.")

        # --- Download Button ---
        st.subheader("Download Comparison Report")
        if st.button("Generate & Prepare DOCX Report", key="compare_download_prep"):
            with st.spinner("Generating DOCX file..."):
                os.makedirs(DEFAULT_OUTPUT_DIR, exist_ok=True)
                primary = primary_ticker.upper() # Use from form submission context
                comps_str = "_vs_" + "_".join(c.strip().upper() for c in competitors_list)
                ts = datetime.now().strftime("%Y%m%d_%H%M%S")
                output_filename = os.path.join(DEFAULT_OUTPUT_DIR, f"{primary}{comps_str}_FY{year}Q{quarter}_Comparison_{ts}.docx")

                success = generate_comparison_docx(primary_ticker, competitors_list, year, quarter, results, output_filename)

                if success:
                    st.session_state.compare_report_path = output_filename
                else:
                    # Error already displayed within generate_comparison_docx
                    pass

        if st.session_state.compare_report_path:
            try:
                with open(st.session_state.compare_report_path, "rb") as file:
                    st.download_button(
                        label="Download Report", data=file,
                        file_name=os.path.basename(st.session_state.compare_report_path),
                        mime="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
                        key="compare_download_final"
                    )
            except FileNotFoundError: st.error("🔴 Error: Generated report file not found.")
            except Exception as e: st.error(f"🔴 Error preparing download: {e}")

    elif st.session_state.compare_error:
        st.error("Comparison failed. See error details below.")
        st.code(st.session_state.compare_error)

# Display error if occurred outside results block
elif st.session_state.compare_error:
     st.error("Comparison failed. See error details below.")
     st.code(st.session_state.compare_error)