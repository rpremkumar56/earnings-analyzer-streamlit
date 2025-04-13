# compare_competitors.py

import argparse
import os
import sys
from datetime import datetime
import pandas as pd
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_PARAGRAPH_ALIGNMENT
from dotenv import load_dotenv
import google.generativeai as genai
import sqlite3
import time
import re

# --- Import Database Manager ---
from database_manager import get_db_connection, get_comparison_data

# --- Gemini Configuration & Helper (Copied from report_generator for self-containment) ---
# (It's better practice to put shared utils like this in a separate file,
# but copying is simpler for this example)
def generate_gemini_content(prompt, retries=3, delay=5):
    """Generates content using the Gemini API with error handling and retries."""
    print(f"  Calling Gemini API for comparison task...")
    MODEL_NAME = "gemini-1.5-flash" # Or potentially a more powerful model for comparison
    SAFETY_SETTINGS = [{"category": c, "threshold": "BLOCK_MEDIUM_AND_ABOVE"} for c in ["HARM_CATEGORY_HARASSMENT", "HARM_CATEGORY_HATE_SPEECH", "HARM_CATEGORY_SEXUALLY_EXPLICIT", "HARM_CATEGORY_DANGEROUS_CONTENT"]]
    GENERATION_CONFIG = {"temperature": 0.7, "top_p": 0.95, "top_k": 64, "max_output_tokens": 4096, "response_mime_type": "text/plain"} # Allow more tokens for comparison
    model = genai.GenerativeModel(model_name=MODEL_NAME, safety_settings=SAFETY_SETTINGS, generation_config=GENERATION_CONFIG)
    # (Rest of the generate_gemini_content function is identical to the one in report_generator.py)
    for attempt in range(retries):
        try:
            response = model.generate_content(prompt)
            if response and response.candidates and response.candidates[0].content.parts:
                 print("  Gemini API call successful.")
                 return response.text.strip()
            else:
                block_reason = response.prompt_feedback.block_reason if response.prompt_feedback else 'Unknown'
                safety_ratings = response.candidates[0].safety_ratings if response.candidates else 'N/A'
                print(f"  Warning: Gemini response blocked or empty. Reason: {block_reason}. Ratings: {safety_ratings}", file=sys.stderr)
                return f"Error: Response blocked or empty (Reason: {block_reason})"
        except Exception as e:
            print(f"  Error calling Gemini API (Attempt {attempt + 1}/{retries}): {e}", file=sys.stderr)
            if attempt < retries - 1:
                print(f"  Retrying in {delay} seconds...")
                time.sleep(delay)
            else:
                print("  Max retries reached. Skipping this request.", file=sys.stderr)
                return f"Error: API call failed after multiple retries ({e})"
    return "Error: Max retries reached."


# --- Word Document Helper Functions (Copied from report_generator) ---
def add_heading(doc, text, level=1):
    doc.add_heading(text, level=level)

def add_paragraph(doc, text, style=None, bold=False, italic=False):
    p = doc.add_paragraph(text, style=style)
    if bold:
        for run in p.runs: run.bold = True
    if italic:
        for run in p.runs: run.italic = True
    return p

def set_col_widths(table, widths):
    for row in table.rows:
        for idx, width in enumerate(widths):
            if idx < len(row.cells): row.cells[idx].width = Inches(width)

# --- Comparison Specific Logic ---

def create_financial_comparison_table(doc, comparison_data, primary_ticker, competitors):
    """Creates the financial comparison table."""
    print("Creating financial comparison table...")
    # Define key metrics and their preferred display name
    metrics_to_compare = {
        "Revenue ($M) Q3": "Revenue ($M)",
        "Operating Margin (%) Q3": "Operating Margin (%)",
        # Add YoY/QoQ Growth if available directly from saved data or calculate here
        # "YoY Revenue Growth (%)": "YoY Revenue Growth (%)", # Example if stored
        "Net Income ($M) Q3 (from Table)": "Net Income ($M)",
        "Net Margin (%) Q3 (calculated from Table)": "Net Margin (%)",
        "Basic EPS ($) Q3": "Basic EPS ($)",
        "Large Deal TCV ($B) Q3": "Large Deal TCV ($B)",
    }
    tickers_in_order = [primary_ticker] + competitors
    table_data = []

    # Header Row for DataFrame
    df_columns = ["Metric"] + tickers_in_order

    # Populate data for DataFrame
    for metric_key, display_name in metrics_to_compare.items():
        row_data = {"Metric": display_name}
        for ticker in tickers_in_order:
            # Use the 'raw' value for display in the table
            raw_value = comparison_data.get(ticker, {}).get("financials", {}).get(metric_key, {}).get('raw', 'N/A')
            row_data[ticker] = raw_value
        table_data.append(row_data)

    if not table_data:
        add_paragraph(doc, "(No comparable financial data found)", italic=True)
        return

    df = pd.DataFrame(table_data, columns=df_columns)

    # Add table to document
    num_rows = df.shape[0] + 1
    num_cols = df.shape[1]
    table = doc.add_table(rows=num_rows, cols=num_cols, style='Table Grid')
    table.autofit = False
    table.allow_autofit = False

    # Set column widths (adjust based on number of competitors)
    metric_width = 2.5
    competitor_width = max(1.2, (6.5 - metric_width) / len(tickers_in_order)) # Distribute remaining width
    col_widths = [metric_width] + [competitor_width] * len(tickers_in_order)
    set_col_widths(table, col_widths)

    # Add header row
    header_cells = table.rows[0].cells
    for j, col_name in enumerate(df.columns):
        p = header_cells[j].paragraphs[0]
        run = p.add_run(col_name)
        run.font.bold = True
        p.alignment = WD_PARAGRAPH_ALIGNMENT.CENTER if j > 0 else WD_PARAGRAPH_ALIGNMENT.LEFT

    # Add data rows
    for i in range(df.shape[0]):
        data_cells = table.rows[i + 1].cells
        for j, col_name in enumerate(df.columns):
            value = str(df.iat[i, j])
            p = data_cells[j].paragraphs[0]
            p.add_run(value)
            # Right-align numeric columns
            if j > 0:
                 p.alignment = WD_PARAGRAPH_ALIGNMENT.RIGHT


def generate_qualitative_comparison(comparison_data, primary_ticker, competitors, theme):
    """Generates comparative text for a specific theme using Gemini."""
    print(f"Generating qualitative comparison for theme: {theme}...")
    tickers_in_order = [primary_ticker] + competitors
    theme_summaries = {}
    found_data = False

    for ticker in tickers_in_order:
        summary_data = comparison_data.get(ticker, {}).get("summaries", {}).get(theme)
        if summary_data and summary_data.get('text'):
            theme_summaries[ticker] = f"Sentiment: {summary_data.get('sentiment', 'N/A')}\nSummary: {summary_data['text']}"
            found_data = True
        else:
             theme_summaries[ticker] = "No specific summary found for this theme."

    if not found_data:
        print(f"  No data found for any company for theme '{theme}'. Skipping comparison.")
        return f"No data available in the database for any selected company regarding '{theme}' for this period."

    # Construct Prompt
    prompt_sections = [f"Compare and contrast the following summaries regarding '{theme}' for the specified companies."]
    prompt_sections.append("Focus on similarities, differences in strategy, key initiatives, reported progress, and overall sentiment where available.\n")

    for ticker in tickers_in_order:
        prompt_sections.append(f"--- {ticker} ---")
        prompt_sections.append(theme_summaries[ticker])
        prompt_sections.append("-" * (len(ticker) + 8) + "\n")

    prompt_sections.append("--- Comparison ---")
    prompt = "\n".join(prompt_sections)

    comparison_text = generate_gemini_content(prompt)

    if comparison_text.startswith("Error:"):
        print(f"  Comparison generation failed for theme {theme}: {comparison_text}")
        return f"Comparison for '{theme}' could not be generated due to an error."
    else:
        print(f"  Comparison generated successfully for theme: {theme}")
        return comparison_text

# --- Main Execution Logic ---
def run_comparison(args):
    """Fetches data and generates the competitor comparison report."""

    db_conn = None
    try:
        # Connect to DB
        db_conn = get_db_connection(args.db_path)

        # Parse competitors
        competitors = [c.strip().upper() for c in args.competitors.split(',')]
        all_tickers = [args.primary_ticker.upper()] + competitors

        # Fetch data
        comparison_data = get_comparison_data(db_conn, all_tickers, args.year, args.quarter)

        # Check if primary company data exists
        if not comparison_data.get(args.primary_ticker.upper(), {}).get("report_id"):
             print(f"Error: Data for primary ticker {args.primary_ticker.upper()} not found in database for FY{args.year} Q{args.quarter}.", file=sys.stderr)
             print("Please run the analysis for the primary ticker first using main.py.")
             return # Exit if primary data is missing

        # --- Generate Report ---
        doc = Document()
        primary_ticker = args.primary_ticker.upper()

        # Title and Metadata
        comp_string = ", ".join(competitors)
        add_heading(doc, f"Competitor Analysis: FY{args.year} Q{args.quarter}", level=0)
        add_paragraph(doc, f"Primary Company: {primary_ticker}")
        add_paragraph(doc, f"Competitors Compared: {comp_string}")
        add_paragraph(doc, f"Generated: {datetime.now().strftime('%Y-%m-%d %H:%M:%S')}")
        doc.add_paragraph()

        # Financial Comparison Table
        add_heading(doc, "Financial Benchmark", level=1)
        create_financial_comparison_table(doc, comparison_data, primary_ticker, competitors)
        doc.add_paragraph()

        # Qualitative Comparison Sections
        add_heading(doc, "Qualitative Comparison (Based on Transcript Summaries)", level=1)
        themes_to_compare = [
            "AI/GenAI",
            "Demand Environment/Pipeline",
            "Margins/Profitability/Costs",
            "Vertical Performance (e.g., BFSI, Retail)", # Example Verticals
            "Geographic Performance"
            # Add/remove themes as needed
        ]

        for theme in themes_to_compare:
             add_heading(doc, theme, level=2)
             comparison_text = generate_qualitative_comparison(comparison_data, primary_ticker, competitors, theme)
             add_paragraph(doc, comparison_text)
             # Check for missing data notes
             if "No data available" in comparison_text or "No specific summary found" in comparison_text:
                 add_paragraph(doc, "(Ensure individual company analyses have been run for this period using main.py)", italic=True)
             doc.add_paragraph()

        # --- Save Document ---
        doc.save(args.output_file)
        print(f"\nComparison report saved successfully to: {args.output_file}")

    except Exception as e:
        print(f"\nAn error occurred during comparison generation: {e}", file=sys.stderr)
        import traceback
        traceback.print_exc()
    finally:
        if db_conn:
            db_conn.close()
            print("Database connection closed.")


if __name__ == "__main__":
    load_dotenv()
    api_key = os.getenv("GOOGLE_API_KEY")
    if not api_key:
         print("CRITICAL ERROR: GOOGLE_API_KEY missing.", file=sys.stderr)
         sys.exit(1)
    try:
        genai.configure(api_key=api_key)
        print("Google Gemini API Key configured.")
    except Exception as e:
        print(f"CRITICAL ERROR: Failed to configure Google Gemini API: {e}", file=sys.stderr)
        sys.exit(1)

    parser = argparse.ArgumentParser(description="Generate a competitor comparison report using stored earnings analysis data.")
    parser.add_argument("--primary_ticker", required=True, help="Ticker symbol for your primary company.")
    parser.add_argument("--competitors", required=True, help="Comma-separated list of competitor ticker symbols (e.g., COMPA,COMPB).")
    parser.add_argument("--year", required=True, type=int, help="Fiscal year to compare.")
    parser.add_argument("--quarter", required=True, type=int, choices=[1, 2, 3, 4], help="Fiscal quarter to compare.")
    parser.add_argument("--output_file", default=None, help="Path to save the output DOCX report. Defaults to a generated name.")
    parser.add_argument("--db_path", default=None, help="Optional path to the SQLite database file.")

    args = parser.parse_args()

    # Generate default output filename if not provided
    if args.output_file is None:
         timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
         competitor_str = "_vs_" + "_".join(c.strip().upper() for c in args.competitors.split(','))
         args.output_file = f"{args.primary_ticker.upper()}{competitor_str}_FY{args.year}Q{args.quarter}_Comparison_{timestamp}.docx"

    run_comparison(args)