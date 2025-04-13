# report_generator.py

import os
import sys
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_PARAGRAPH_ALIGNMENT
from docx.enum.style import WD_STYLE_TYPE
import pandas as pd
import google.generativeai as genai
import time
import re # For parsing comparison percentages

# --- Gemini Configuration & Helper (Keep as before) ---

def generate_gemini_content(prompt, retries=3, delay=5):
    """Generates content using the Gemini API with error handling and retries."""
    print(f"  Calling Gemini API for report generation task...")
    MODEL_NAME = "gemini-1.5-flash"
    SAFETY_SETTINGS = [{"category": c, "threshold": "BLOCK_MEDIUM_AND_ABOVE"} for c in ["HARM_CATEGORY_HARASSMENT", "HARM_CATEGORY_HATE_SPEECH", "HARM_CATEGORY_SEXUALLY_EXPLICIT", "HARM_CATEGORY_DANGEROUS_CONTENT"]]
    GENERATION_CONFIG = {"temperature": 0.6, "top_p": 0.95, "top_k": 64, "max_output_tokens": 2048, "response_mime_type": "text/plain"}
    model = genai.GenerativeModel(model_name=MODEL_NAME, safety_settings=SAFETY_SETTINGS, generation_config=GENERATION_CONFIG)
    for attempt in range(retries):
        try:
            response = model.generate_content(prompt)
            if response and response.candidates and response.candidates[0].content.parts:
                print("  Gemini API call successful.")
                return response.text.strip()
            else:
                block_reason = response.prompt_feedback.block_reason if response.prompt_feedback else 'Unknown'
                safety_ratings = response.candidates[0].safety_ratings if response.candidates else 'N/A'
                print(f"  Warning: Gemini response blocked or empty. Reason: {block_reason}. Ratings: {safety_ratings}", file=sys.stderr)
                return f"Error: Response blocked or empty (Reason: {block_reason})"
        except Exception as e:
            print(f"  Error calling Gemini API (Attempt {attempt + 1}/{retries}): {e}", file=sys.stderr)
            if attempt < retries - 1:
                print(f"  Retrying in {delay} seconds...")
                time.sleep(delay)
            else:
                print("  Max retries reached. Skipping this request.", file=sys.stderr)
                return f"Error: API call failed after multiple retries ({e})"
    return "Error: Max retries reached."

# --- Helper Functions for Word Document (Keep add_heading, add_paragraph, add_bullet_point) ---

def add_heading(doc, text, level=1):
    doc.add_heading(text, level=level)

def add_paragraph(doc, text, style=None, bold=False, italic=False):
    p = doc.add_paragraph(text, style=style)
    if bold:
        for run in p.runs: run.bold = True
    if italic:
        for run in p.runs: run.italic = True
    return p

def add_bullet_point(doc, text, indent_level=1):
    # Using List Bullet style; actual indentation depends on template/defaults
    p = doc.add_paragraph(style='List Bullet')
    p.add_run(text) # Add text to the paragraph with list style

def set_col_widths(table, widths):
    """Sets column widths (in inches)."""
    for row in table.rows:
        for idx, width in enumerate(widths):
            if idx < len(row.cells): # Check if cell index exists
                 row.cells[idx].width = Inches(width)

def add_financials_table(doc, financial_summary, comparisons):
    """Creates a formatted table for key financials with comparisons."""
    print("Creating financial highlights table...")
    # Define which metrics to include in the table and their display order
    metrics_in_table = [
        "Revenue ($M) Q3",
        "Operating Margin (%) Q3",
        "Net Margin (%) Q3 (calculated from Table)",
        "Net Income ($M) Q3 (from Table)",
        "Basic EPS ($) Q3",
        "Large Deal TCV ($B) Q3",
        "Free Cash Flow ($M) Q3",
    ]

    # Filter and sort the summary data
    table_data = []
    for metric in metrics_in_table:
        if metric in financial_summary:
             table_data.append({
                 "Metric": metric.replace(" Q3", "").replace(" ($M)", " $M").replace(" ($B)", " $B").replace(" (%)", " %").replace("(calculated from Table)","").replace("(from Table)","").strip(), # Clean up name
                 "Current Value": financial_summary.get(metric, "N/A"),
                 "YoY Change": comparisons.get("yoy_pct_change", {}).get(metric, "N/A"),
                 "QoQ Change": comparisons.get("qoq_pct_change", {}).get(metric, "N/A")
             })

    if not table_data:
        add_paragraph(doc, "(No key financial data available for table)", italic=True)
        return

    df = pd.DataFrame(table_data)

    # Add table to document
    num_rows = df.shape[0] + 1
    num_cols = df.shape[1]
    table = doc.add_table(rows=num_rows, cols=num_cols, style='Table Grid')
    table.autofit = False # Disable autofit to set widths manually
    table.allow_autofit = False

    # Set column widths (adjust as needed)
    col_widths = [2.5, 1.2, 1.2, 1.2] # Inches for Metric, Current, YoY, QoQ
    set_col_widths(table, col_widths)


    # Add header row
    header_cells = table.rows[0].cells
    for j, col_name in enumerate(df.columns):
        p = header_cells[j].paragraphs[0]
        run = p.add_run(col_name)
        run.font.bold = True
        p.alignment = WD_PARAGRAPH_ALIGNMENT.CENTER

    # Define colors for changes
    green_color = RGBColor(0x00, 0x80, 0x00) # Dark Green
    red_color = RGBColor(0xFF, 0x00, 0x00)   # Red

    # Add data rows with color coding for changes
    for i in range(df.shape[0]):
        data_cells = table.rows[i + 1].cells
        for j, col_name in enumerate(df.columns):
            value = str(df.iat[i, j])
            p = data_cells[j].paragraphs[0]
            run = p.add_run(value)

            # Apply color to YoY and QoQ change columns
            if col_name in ["YoY Change", "QoQ Change"] and value not in ["N/A", "Error", "Inf", "0.0%"]:
                try:
                    # Extract numeric part for comparison
                    num_val_match = re.match(r"^(-?[\d\.]+)", value)
                    if num_val_match:
                        num_val = float(num_val_match.group(1))
                        if num_val > 0:
                            run.font.color.rgb = green_color
                            run.text = f"+{value}" # Add explicit plus sign
                        elif num_val < 0:
                            run.font.color.rgb = red_color
                        # else: keep default color for 0.0%
                except ValueError:
                    pass # Ignore if conversion fails

            # Align values (optional)
            if col_name != "Metric":
                 p.alignment = WD_PARAGRAPH_ALIGNMENT.RIGHT


# --- Updated Executive Summary Generation ---

def generate_executive_summary(report_bundle):
    """Uses Gemini to generate summary leveraging historical context."""
    print("Generating Executive Summary using Gemini...")
    current_data = report_bundle['current_data']
    comparisons = report_bundle['comparisons']
    # historical_metrics = report_bundle['historical_metrics'] # Raw historical if needed

    metadata = current_data.get('metadata', {})
    pdf_data = current_data.get('pdf_data', {})
    transcript_data = current_data.get('transcript_data', {})
    financials = pdf_data.get('financial_summary', {})
    pdf_outlook = pdf_data.get('extracted_sections', {}).get('Outlook', '')
    transcript_demand_summary = transcript_data.get('commentary_summaries', {}).get('Demand Environment/Pipeline', '')

    # --- Key Data Points ---
    ticker = metadata.get('ticker', 'Company')
    fy = metadata.get('fiscal_year', 'N/A')
    fq = metadata.get('fiscal_quarter', 'N/A')
    revenue = financials.get('Revenue ($M) Q3', 'N/A')
    op_margin = financials.get('Operating Margin (%) Q3', 'N/A')
    eps = financials.get('Basic EPS ($) Q3', 'N/A')
    tcv = financials.get('Large Deal TCV ($B) Q3', 'N/A')
    fy_rev_guidance = financials.get('FY25 Revenue Guidance (%)', 'N/A')
    fy_margin_guidance = financials.get('FY25 Operating Margin Guidance (%)', 'N/A')

    # --- Key Comparison Points ---
    rev_yoy = comparisons.get('yoy_pct_change', {}).get('Revenue ($M) Q3', 'N/A')
    op_margin_yoy = comparisons.get('yoy_pct_change', {}).get('Operating Margin (%) Q3', 'N/A')
    eps_yoy = comparisons.get('yoy_pct_change', {}).get('Basic EPS ($) Q3', 'N/A')
    # Add QoQ if desired
    rev_qoq = comparisons.get('qoq_pct_change', {}).get('Revenue ($M) Q3', 'N/A')

    # --- Limit snippet length ---
    pdf_outlook_snippet = pdf_outlook[:800]
    transcript_demand_snippet = transcript_demand_summary[:800]

    # --- Construct Updated Prompt ---
    prompt = f"""
    Generate a concise executive summary (1-2 short paragraphs) for an earnings analysis report for {ticker} (FY{fy} Q{fq}).
    Focus on the quarter's performance relative to expectations (implied by YoY/QoQ changes) and the future outlook based on guidance and commentary.

    **Current Quarter Performance (FY{fy} Q{fq}):**
    *   Revenue ($M): {revenue} (YoY: {rev_yoy}, QoQ: {rev_qoq})
    *   Operating Margin (%): {op_margin} (YoY: {op_margin_yoy})
    *   Basic EPS ($): {eps} (YoY: {eps_yoy})
    *   Large Deal TCV ($B): {tcv}

    **Guidance (Full Year):**
    *   Revenue Guidance (%): {fy_rev_guidance}
    *   Operating Margin Guidance (%): {fy_margin_guidance}
    *   (Note if guidance appears revised based on commentary or data)

    **Commentary Snippets:**
    *   Official Outlook (from PDF): "{pdf_outlook_snippet}..."
    *   Transcript Discussion (Demand/Pipeline): "{transcript_demand_snippet}..."

    **Task:** Synthesize this information into a brief summary. Highlight key performance trends (YoY growth/decline in revenue/EPS/margins), mention the guidance stance (raised/lowered/maintained if apparent), and capture the essence of the management outlook/commentary.
    """

    summary_text = generate_gemini_content(prompt)

    if summary_text.startswith("Error:"):
        print(f"  Executive Summary generation failed: {summary_text}")
        return "Executive summary could not be generated due to an error."
    else:
        print("  Executive Summary generated successfully.")
        return summary_text


# --- Updated Core Report Generation Logic ---

def create_word_report(report_bundle, output_filename):
    """
    Generates the full analysis report in DOCX format using bundled data.
    """
    print(f"--- Generating DOCX Report: {output_filename} ---")
    os.makedirs(os.path.dirname(output_filename), exist_ok=True)

    try:
        doc = Document()

        # --- Extract data from bundle ---
        current_data = report_bundle['current_data']
        historical_metrics = report_bundle['historical_metrics'] # Raw historical data
        comparisons = report_bundle['comparisons'] # Calculated % changes

        metadata = current_data.get('metadata', {})
        pdf_data = current_data.get('pdf_data', {})
        transcript_data = current_data.get('transcript_data', {})
        financial_summary = pdf_data.get('financial_summary', {}) # Current raw values

        # --- Title and Metadata (Updated) ---
        ticker = metadata.get('ticker', 'N/A')
        fy = metadata.get('fiscal_year', 'N/A')
        fq = metadata.get('fiscal_quarter', 'N/A')
        report_date = metadata.get('report_date', 'N/A')
        add_heading(doc, f"{ticker} - FY{fy} Q{fq} Earnings Analysis", level=0)
        add_paragraph(doc, f"Period Ending: {report_date}")
        add_paragraph(doc, f"PDF Source: {metadata.get('pdf_filename', 'N/A')}")
        add_paragraph(doc, f"Transcript Source: {metadata.get('transcript_filename', 'N/A')}")
        add_paragraph(doc, f"Analysis Date: {metadata.get('analysis_date', 'N/A')}")
        doc.add_paragraph()

        # --- Executive Summary (Generated by Gemini, uses bundle) ---
        add_heading(doc, "Executive Summary", level=1)
        summary_text = generate_executive_summary(report_bundle) # Pass the whole bundle
        add_paragraph(doc, summary_text)
        doc.add_paragraph()

        # --- Key Financial Highlights (New Table Format) ---
        add_heading(doc, "Key Financial Highlights", level=1)
        # Pass current summary and calculated comparisons to the table function
        add_financials_table(doc, financial_summary, comparisons)
        doc.add_paragraph()

        # --- Guidance ---
        add_heading(doc, "Full Year Guidance", level=2)
        guidance_found = False
        # Extract guidance from current financial summary
        rev_guidance = financial_summary.get('FY25 Revenue Guidance (%)', 'Not Found') # Adjust key if year changes
        margin_guidance = financial_summary.get('FY25 Operating Margin Guidance (%)', 'Not Found')

        if rev_guidance != 'Not Found':
            add_bullet_point(doc, f"Revenue Growth: {rev_guidance}")
            guidance_found = True
        if margin_guidance != 'Not Found':
            add_bullet_point(doc, f"Operating Margin: {margin_guidance}")
            guidance_found = True

        # Optional: Compare to prior guidance if available in historical_metrics
        prior_rev_guidance = historical_metrics.get("qoq", {}).get('FY25 Revenue Guidance (%)', {}).get('raw', 'N/A') # Example key
        if prior_rev_guidance != 'N/A' and prior_rev_guidance != rev_guidance:
             add_paragraph(doc, f"(Prior Quarter Revenue Guidance: {prior_rev_guidance})", italic=True)
             guidance_found = True
        # Add similar logic for margin guidance comparison if needed

        if not guidance_found:
             add_paragraph(doc, "Guidance figures not found or unchanged.", italic=True)
        doc.add_paragraph()

        # --- PDF Analysis (Accessing data from current_data) ---
        add_heading(doc, "PDF Analysis Details", level=1)
        extracted_sections = pdf_data.get('extracted_sections', {})
        if extracted_sections:
             add_heading(doc, "Key Sections from Press Release", level=2)
             # Prioritize Outlook and Management Commentary
             for section_name in ["Outlook", "Management Commentary", "Financial Highlights"]:
                  if section_name in extracted_sections:
                       add_heading(doc, section_name, level=3)
                       section_text = extracted_sections[section_name]
                       add_paragraph(doc, section_text[:1500] + ('...' if len(section_text) > 1500 else ''))
                       doc.add_paragraph()
             # Add other sections if needed
        # Table Identification Status
        add_heading(doc, "Table Identification Status", level=2)
        identified_tables = pdf_data.get("identified_tables")
        if identified_tables:
            for table_key, table_obj in identified_tables.items():
                status = "Found" if table_obj is not None else "Not Found"
                add_bullet_point(doc, f"{table_key.replace('_', ' ').title()}: {status}")
        else:
            add_paragraph(doc, "No tables were specifically identified.")
        # Optional: Embed Identified Income Statement Table (Keep as before)
        if identified_tables and identified_tables.get("income_statement") is not None:
            add_heading(doc, "Identified Income Statement (Preview)", level=2)
            # (Code to add table preview remains the same)
            try:
                preview_df = identified_tables["income_statement"].head(15).iloc[:, :5]
                # Need the helper function here if not imported globally
                # Reuse add_table_from_dataframe from above if suitable or simplify
                temp_table = doc.add_table(preview_df.shape[0] + 1, preview_df.shape[1], style='Table Grid')
                temp_table.autofit = True
                for j, col_name in enumerate(preview_df.columns): temp_table.cell(0, j).text = str(col_name)
                for i in range(preview_df.shape[0]):
                    for j in range(preview_df.shape[1]):
                        temp_table.cell(i + 1, j).text = str(preview_df.iat[i, j])
            except Exception as e:
                 add_paragraph(doc, f"(Error generating table preview: {e})", italic=True)
            doc.add_paragraph()


        # --- Transcript Analysis (Accessing data from current_data) ---
        add_heading(doc, "Earnings Call Transcript Analysis", level=1)
        if transcript_data:
            if transcript_data.get("error"):
                add_paragraph(doc, f"Transcript Processing Error: {transcript_data['error']}", bold=True)
            else:
                # Executive Quotes (Keep as before)
                add_heading(doc, "Executive Quotes (from Press Release)", level=2)
                # (Code remains the same)
                executive_quotes = transcript_data.get("executive_quotes")
                if executive_quotes:
                    for role, quote in executive_quotes.items():
                        add_paragraph(doc, f"[{role}]", bold=True)
                        add_paragraph(doc, quote)
                        doc.add_paragraph()
                else:
                    add_paragraph(doc, "No CEO/CFO quotes extracted.", italic=True)

                # Thematic Summaries (Keep as before, parsing sentiment)
                add_heading(doc, "Commentary Summaries by Theme (from Transcript)", level=2)
                # (Code remains the same)
                commentary_summaries = transcript_data.get("commentary_summaries")
                if commentary_summaries:
                    for theme, summary_text in commentary_summaries.items():
                        add_heading(doc, theme, level=3)
                        lines = summary_text.splitlines()
                        sentiment, summary_body = "", []
                        for line in lines:
                             if line.strip().startswith("Sentiment:"): sentiment = line.strip()
                             elif line.strip().startswith("Summary:"): summary_body.append(line.replace("Summary:", "").strip())
                             else: summary_body.append(line.strip())
                        if sentiment: add_paragraph(doc, sentiment, bold=True)
                        summary_content = "\n".join(s for s in summary_body if s).strip()
                        if summary_content:
                             if summary_content.startswith("*") or summary_content.startswith("-"):
                                 for item in summary_content.split('\n'):
                                     if item.strip(): add_bullet_point(doc, item.strip().lstrip('*- '))
                             else: add_paragraph(doc, summary_content)
                        else: add_paragraph(doc,"(No summary content)", italic=True)
                        doc.add_paragraph()
                else:
                    add_paragraph(doc, "No thematic summaries generated.", italic=True)

                # Keywords (Keep as before)
                add_heading(doc, "Top Keywords (from Transcript)", level=2)
                # (Code remains the same)
                keywords = transcript_data.get("keywords")
                if keywords:
                    keyword_list = [kw[0] for kw in keywords[:20]]
                    add_paragraph(doc, ", ".join(keyword_list))
                else:
                    add_paragraph(doc, "No keywords extracted.", italic=True)
        else:
             add_paragraph(doc, "No transcript data processed.", italic=True)

        # --- Save Document ---
        doc.save(output_filename)
        print(f"--- Report saved successfully to {output_filename} ---")
        return True

    except Exception as e:
        print(f"--- Error generating Word report: {e} ---", file=sys.stderr)
        import traceback
        traceback.print_exc()
        return False

# Example usage (Update to use report_bundle structure)
if __name__ == "__main__":
    print("Testing report_generator.py with historical context...")
    # Create dummy data structure similar to what main.py provides
    dummy_report_bundle = {
         "current_data": {
             "metadata": {"ticker": "DUMMY", "fiscal_year": 2024, "fiscal_quarter": 3, "report_date": "2024-09-30", "pdf_filename": "dummy.pdf", "transcript_filename": "dummy.docx", "analysis_date": "2023-10-27 10:00:00"},
             "pdf_data": {
                 "financial_summary": {"Revenue ($M) Q3": "1050", "Operating Margin (%) Q3": "21.0", "Basic EPS ($) Q3": "0.22", "Large Deal TCV ($B) Q3": "3.1", "FY25 Revenue Guidance (%)": "5%-6%", "FY25 Operating Margin Guidance (%)": "21%-22%"},
                 "extracted_sections": {"Outlook": "Positive outlook continues.", "Management Commentary": "Solid execution."},
                 "identified_tables": {"income_statement": pd.DataFrame({'A': ['Rev', 'Profit'], 'B': [1050, 220]})}
             },
             "transcript_data": {"executive_quotes": {"CEO": '"Good progress" - CEO'}, "commentary_summaries": {"Demand": "Sentiment: Positive\nSummary: Stronger than last quarter."}, "keywords": [("cloud", 1), ("AI", 1)]}
         },
         "historical_metrics": {
             "qoq": {"Revenue ($M) Q3": {"value": 1000.0, "raw": "1000"}, "Operating Margin (%) Q3": {"value": 20.5, "raw": "20.5"}},
             "yoy": {"Revenue ($M) Q3": {"value": 950.0, "raw": "950"}, "Operating Margin (%) Q3": {"value": 21.5, "raw": "21.5"}}
         },
         "comparisons": {
             "qoq_pct_change": {"Revenue ($M) Q3": "+5.0%", "Operating Margin (%) Q3": "+2.4%"}, # Example calculated values
             "yoy_pct_change": {"Revenue ($M) Q3": "+10.5%", "Operating Margin (%) Q3": "-2.3%"}
         }
    }
    # Configure API Key for the test run
    from dotenv import load_dotenv
    load_dotenv()
    if not os.getenv("GOOGLE_API_KEY"):
        print("CRITICAL: GOOGLE_API_KEY not found in .env for testing report generator.")
    else:
        genai.configure(api_key=os.getenv("GOOGLE_API_KEY"))
        test_output_filename = "test_report_output_with_history.docx"
        success = create_word_report(dummy_report_bundle, test_output_filename)
        if success:
            print(f"Test report generated: {test_output_filename}")
        else:
            print("Test report generation failed.")